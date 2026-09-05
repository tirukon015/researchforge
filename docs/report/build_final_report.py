"""Build ResearchForge_Final_Project_Report.docx from FINAL_REPORT_SOURCE.md.

    python docs/report/build_final_report.py

The Markdown file is the source; this renders it as a Microsoft Word document
with the formatting the assignment requires: Times New Roman 11 pt, 1.5 line
spacing, justified body text, real Word heading styles, a live Table of
Contents, page numbers, and figures and tables with captions.

FIELDS RATHER THAN TYPED LISTS
------------------------------
The contents page is a real Word TOC field, so it paginates itself. The lists
of figures and tables cannot be TOC fields here: Word builds those from SEQ
fields, which number sequentially in document order, and this report's figure
numbers are fixed by the assignment brief (the screenshots are Figures 4.4 to
4.11 even though they appear before the diagrams numbered 4.1 to 4.3). Each
caption is therefore bookmarked, and each list entry carries a PAGEREF field
pointing at that bookmark, so the page numbers are still computed by Word
rather than typed.

Every field is marked dirty, so Word offers to update them when the file is
opened; Ctrl+A then F9 updates them all at once.

DIRECTIVES UNDERSTOOD
---------------------
    @FRONTMATTER              the title page
    @TOC @LISTOFFIGURES @LISTOFTABLES
    @PAGEBREAK
    @FIG|image|Figure 4.4|caption
    @TABLE|Table 4.1|caption   (followed by a Markdown table)
    @CODE ... @ENDCODE
"""

from __future__ import annotations

import pathlib
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "submission"
SHOTS = HERE / "screenshots"
DIAGRAMS = HERE / "diagrams"

FONT = "Times New Roman"
BODY_PT = 11
BLACK = RGBColor(0, 0, 0)

# Text column on A4 with 2.54 cm margins.
COL_W = 6.27
MAX_FIG_H = 7.6


# --------------------------------------------------------------- field helpers

def _field(paragraph, instr: str, placeholder: str = ""):
    """Insert a real Word field, marked dirty so Word refreshes it on open."""
    r1 = paragraph.add_run()._r
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    fld.set(qn("w:dirty"), "true")
    r1.append(fld)

    r2 = paragraph.add_run()._r
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2.append(it)

    r3 = paragraph.add_run()._r
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3.append(sep)

    if placeholder:
        run = paragraph.add_run(placeholder)
        run.font.name = FONT
        run.font.size = Pt(BODY_PT)

    r4 = paragraph.add_run()._r
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r4.append(end)


def _bookmark(paragraph, name: str, bid: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _tab_stop_right(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), str(int(COL_W * 1440)))
    tabs.append(tab)
    pPr.append(tabs)


# ------------------------------------------------------------- inline markdown

_LEVELS = [
    (re.compile(r"(\*\*.+?\*\*)", re.S), lambda s: (s[2:-2], "bold")),
    (re.compile(r"((?<!\*)\*[^*\n]+?\*(?!\*))"), lambda s: (s[1:-1], "italic")),
    (re.compile(r"(`[^`\n]+?`)"), lambda s: (s[1:-1], "mono")),
]


def add_inline(paragraph, text, *, size=BODY_PT, bold=False, italic=False,
               mono=False, _level=0):
    if _level < len(_LEVELS):
        pattern, unwrap = _LEVELS[_level]
        for piece in pattern.split(text):
            if not piece:
                continue
            if pattern.fullmatch(piece):
                inner, flag = unwrap(piece)
                add_inline(paragraph, inner, size=size,
                           bold=bold or flag == "bold",
                           italic=italic or flag == "italic",
                           mono=mono or flag == "mono", _level=_level + 1)
            else:
                add_inline(paragraph, piece, size=size, bold=bold, italic=italic,
                           mono=mono, _level=_level + 1)
        return
    run = paragraph.add_run(text)
    run.bold, run.italic = bold, italic
    run.font.name = "Consolas" if mono else FONT
    run.font.size = Pt(size - 1 if mono else size)
    run.font.color.rgb = BLACK


# --------------------------------------------------------------------- document

def new_document() -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_PT)
    st.font.color.rgb = BLACK
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12),
                       ("Caption", 10)):
        s = doc.styles[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.color.rgb = BLACK
        s.font.bold = name != "Caption"
        s.font.italic = name == "Caption"

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.54)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.different_first_page_header_footer = True   # keeps the title page clean
    return doc


def page_numbers(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        _field(p, " PAGE ", "1")
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(10)


def title_page(doc):
    for _ in range(3):
        doc.add_paragraph()
    for text, size, bold in (
        ("AI RESEARCH PAPER ASSISTANT", 22, True),
        ("ResearchForge", 18, False),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(text)
        r.bold, r.font.size, r.font.name = bold, Pt(size), FONT
        r.font.color.rgb = BLACK

    doc.add_paragraph()
    rows = [
        ("Course", "BIT4543 Artificial Intelligence"),
        ("Project area", "AI Topic: Document Intelligence"),
        ("Group members", "[GROUP MEMBER 1 - FULL NAME & STUDENT ID]\n"
                          "[GROUP MEMBER 2 - FULL NAME & STUDENT ID]\n"
                          "[GROUP MEMBER 3 - FULL NAME & STUDENT ID]"),
        ("Lecturer", "Encik Azuan Nazeer"),
        ("Submission date", "[SUBMISSION DATE]"),
        ("Live system", "https://researchforge.rukon.dev"),
        ("Repository", "https://github.com/tirukon015/researchforge"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in rows:
        cells = table.add_row().cells
        for idx, val in enumerate((k, v)):
            cells[idx].text = ""
            para = cells[idx].paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            r = para.add_run(val)
            r.font.name, r.font.size, r.bold = FONT, Pt(11), (idx == 0)
            r.font.color.rgb = BLACK
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def heading(doc, text, level):
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for r in h.runs:
        r.text = ""
    add_inline(h, text, size={1: 16, 2: 13, 3: 12}[level], bold=True)
    return h


def add_image(doc, name, label, caption, state):
    path = SHOTS / f"{name}.png"
    if not path.exists():
        path = DIAGRAMS / f"{name}.png"
    if not path.exists():
        raise SystemExit(f"MISSING IMAGE: {name}")

    from PIL import Image
    w, h = Image.open(path).size
    scale = min(COL_W / w, MAX_FIG_H / h)
    width = w * scale

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(str(path), width=Inches(width))

    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = cap.add_run(f"{label}: {caption}")
    run.font.name, run.font.size, run.italic = FONT, Pt(10), True
    run.font.color.rgb = BLACK
    state["bid"] += 1
    bm = f"_Fig{state['bid']}"
    _bookmark(cap, bm, state["bid"])
    state["figures"].append((label, caption, bm))


def add_table_block(doc, label, caption, rows, state):
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.space_before = Pt(10)
    cap.paragraph_format.space_after = Pt(3)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = cap.add_run(f"{label}: {caption}")
    run.font.name, run.font.size, run.italic = FONT, Pt(10), True
    run.font.color.rgb = BLACK
    state["bid"] += 1
    bm = f"_Tbl{state['bid']}"
    _bookmark(cap, bm, state["bid"])
    state["tables"].append((label, caption, bm))

    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(cols):
            cells[j].text = ""
            para = cells[j].paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_inline(para, row[j] if j < len(row) else "", size=9, bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("\n".join(lines))
    run.font.name, run.font.size = "Consolas", Pt(9)
    run.font.color.rgb = BLACK


def listing(doc, title, entries):
    heading(doc, title, 1)
    for label, caption, bm in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _tab_stop_right(p)
        r = p.add_run(f"{label}: {caption}\t")
        r.font.name, r.font.size = FONT, Pt(11)
        r.font.color.rgb = BLACK
        _field(p, f" PAGEREF {bm} \\h ", "1")


# ------------------------------------------------------------------- rendering

def build():
    src = (HERE / "FINAL_REPORT_SOURCE.md").read_text(encoding="utf-8")
    doc = new_document()
    state = {"figures": [], "tables": [], "bid": 100}

    # Pass 1 renders the body into a temporary document so that the lists of
    # figures and tables are known before they are written; pass 2 renders the
    # real document with those lists in place.
    for final in (False, True):
        doc = new_document()
        state = {"figures": [], "tables": [], "bid": 100}
        render(doc, src, state, final)
        if not final:
            figures, tables = list(state["figures"]), list(state["tables"])
            FIG_CACHE.clear(); FIG_CACHE.extend(figures)
            TBL_CACHE.clear(); TBL_CACHE.extend(tables)

    page_numbers(doc)
    OUT.mkdir(exist_ok=True)
    path = OUT / "ResearchForge_Final_Project_Report.docx"
    doc.save(path)
    return path, state


FIG_CACHE: list = []
TBL_CACHE: list = []


def render(doc, src, state, final):
    lines = src.splitlines()
    i = 0
    pending: list[str] = []
    pending_table = None

    def flush_para():
        nonlocal pending
        if not pending:
            return
        text = " ".join(x.strip() for x in pending).strip()
        pending = []
        if not text:
            return
        if text.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, text[2:])
        elif re.match(r"^\d+\.\s", text):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, re.sub(r"^\d+\.\s+", "", text))
        else:
            p = doc.add_paragraph()
            add_inline(p, text)

    def flush_table():
        nonlocal pending_table
        if pending_table:
            label, caption, rows = pending_table
            add_table_block(doc, label, caption, rows, state)
            pending_table = None

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.startswith("@"):
            flush_para(); flush_table()
            d = line[1:]
            if d == "FRONTMATTER":
                title_page(doc)
            elif d == "PAGEBREAK":
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif d == "TOC":
                heading(doc, "TABLE OF CONTENTS", 1)
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                _field(p, ' TOC \\o "1-3" \\h \\z \\u ',
                       "Right-click and choose Update Field to build the contents.")
            elif d == "LISTOFFIGURES":
                listing(doc, "LIST OF FIGURES", FIG_CACHE if final else [])
            elif d == "LISTOFTABLES":
                listing(doc, "LIST OF TABLES", TBL_CACHE if final else [])
            elif d.startswith("FIG|"):
                _, name, label, caption = d.split("|", 3)
                add_image(doc, name, label, caption, state)
            elif d.startswith("TABLE|"):
                _, label, caption = d.split("|", 2)
                rows = []
                j = i + 1
                while j < len(lines) and lines[j].startswith("|"):
                    cells = [c.strip() for c in lines[j].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells if c):
                        rows.append(cells)
                    j += 1
                pending_table = (label, caption, rows)
                i = j - 1
                flush_table()
            elif d == "CODE":
                buf = []
                j = i + 1
                while j < len(lines) and lines[j].strip() != "@ENDCODE":
                    buf.append(lines[j])
                    j += 1
                add_code(doc, buf)
                i = j
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            flush_para(); flush_table()
            heading(doc, m.group(2).strip(), len(m.group(1)))
            i += 1
            continue

        if not line.strip():
            flush_para()
            i += 1
            continue

        if line.startswith(("- ", "* ")) or re.match(r"^\d+\.\s", line):
            flush_para()
            pending = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].startswith(
                    ("@", "#", "- ", "* ", "|")) and not re.match(r"^\d+\.\s", lines[i]):
                pending.append(lines[i])
                i += 1
            flush_para()
            continue

        pending.append(line)
        i += 1

    flush_para(); flush_table()


if __name__ == "__main__":
    path, state = build()
    print(f"  {path.name}")
    print(f"  {path.stat().st_size:,} bytes")
    print(f"  figures {len(state['figures'])}   tables {len(state['tables'])}")
