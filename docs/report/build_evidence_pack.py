"""Build ResearchForge_Final_Evidence_Pack.docx.

    python docs/report/build_evidence_pack.py

Renders FINAL_EVIDENCE_PACK_SOURCE.md using the same rendering engine as the
final report, so the two documents match in typography and table styling. The
pack carries no figures, so it needs no figure handling; it is tables and
reference text throughout.

Cross-references point at the FINAL report's section numbering, which differs
from the earlier draft's - the sections were renumbered when the report was
restructured to the lecturer's required layout.
"""

from __future__ import annotations

import pathlib
import re
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.shared import Cm, Pt

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build_final_report as R  # noqa: E402

OUT = HERE / "submission"


def title_page(doc):
    for _ in range(4):
        doc.add_paragraph()
    for text, size, bold in (("EVIDENCE PACK", 22, True),
                             ("ResearchForge - AI Research Paper Assistant", 15, False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(text)
        r.bold, r.font.size, r.font.name = bold, Pt(size), R.FONT
        r.font.color.rgb = R.BLACK

    doc.add_paragraph()
    rows = [
        ("Course", "BIT4543 Artificial Intelligence"),
        ("Document", "Supporting evidence for the final project report"),
        ("Group members", "[GROUP MEMBER 1 - FULL NAME & STUDENT ID]\n"
                          "[GROUP MEMBER 2 - FULL NAME & STUDENT ID]\n"
                          "[GROUP MEMBER 3 - FULL NAME & STUDENT ID]"),
        ("Lecturer", "Encik Azuan Nazeer"),
        ("Submission date", "[SUBMISSION DATE]"),
        ("Live system", "https://researchforge.rukon.dev"),
        ("Repository", "https://github.com/tirukon015/researchforge"),
    ]
    from docx.enum.table import WD_TABLE_ALIGNMENT
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in rows:
        cells = table.add_row().cells
        for idx, val in enumerate((k, v)):
            cells[idx].text = ""
            para = cells[idx].paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = para.add_run(val)
            run.font.name, run.font.size, run.bold = R.FONT, Pt(11), (idx == 0)
            run.font.color.rgb = R.BLACK
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build():
    src = (HERE / "FINAL_EVIDENCE_PACK_SOURCE.md").read_text(encoding="utf-8")
    doc = R.new_document()
    state = {"figures": [], "tables": [], "bid": 500}

    # The pack's own title page replaces the report's.
    src = src.replace("@FRONTMATTER\n", "", 1)
    title_page(doc)

    R.render(doc, src, state, True)
    R.page_numbers(doc)

    OUT.mkdir(exist_ok=True)
    path = OUT / "ResearchForge_Final_Evidence_Pack.docx"
    doc.save(path)
    return path, state


if __name__ == "__main__":
    p, st = build()
    print(f"  {p.name}")
    print(f"  {p.stat().st_size:,} bytes")
    print(f"  tables: {len(st['tables'])}")
