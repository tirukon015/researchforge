"""Build ResearchForge_Final_Presentation.pptx on the supplied Slidesgo template.

    python docs/report/build_presentation.py

The template is opened as the base file, so its theme, master, gradient
background, decorative artwork and fonts are inherited rather than imitated.
Its 48 example slides are removed and fifteen ResearchForge slides are built on
its own layouts.

WHY THE DIAGRAMS ARE DRAWN RATHER THAN IMPORTED
-----------------------------------------------
The report's Mermaid diagrams are black-on-white and would sit on this deck's
dark gradient like pasted paper. The same four flows are therefore redrawn here
as native PowerPoint shapes in the template's own palette, which also keeps them
editable.

Every figure on the slides comes from the verified project evidence used in the
final report; nothing is estimated.
"""

from __future__ import annotations

import copy
import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "submission"
SHOTS = HERE / "screenshots"
TEMPLATE = pathlib.Path(
    r"C:\Users\TOUHIDUL ISLAM RUKON\Downloads"
    r"\Cyber-Futuristic AI Technology Thesis Defense by Slidesgo.pptx"
)

# Template palette and typography.
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x51, 0xAF, 0xDB)
TEAL = RGBColor(0x3B, 0x87, 0x94)
DEEP = RGBColor(0x03, 0x11, 0x26)
NAVY = RGBColor(0x10, 0x35, 0x5F)
MUTED = RGBColor(0xB8, 0xCE, 0xE4)
DISPLAY = "Audiowide"
BODY = "Karla"

# Layout indices on master 0.
L_TITLE, L_TITLE_ONLY = 0, 4

# Content band beneath the title frame.
X0, X1 = 0.78, 9.22
Y0, Y1 = 1.42, 5.30
W = X1 - X0


def _in(v):
    return Emu(int(v * 914400))


# ------------------------------------------------------------------ scaffolding

def clear_slides(prs):
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)


def title_frame(slide):
    """The template's snipped-corner outline that sits behind every title."""
    s = slide.shapes.add_shape(MSO_SHAPE.SNIP_2_DIAG_RECTANGLE,
                               _in(0.78), _in(0.56), _in(8.43), _in(0.69))
    s.fill.background()
    s.line.color.rgb = WHITE
    s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def add_slide(prs, title, layout=L_TITLE_ONLY):
    slide = prs.slides.add_slide(prs.slide_masters[0].slide_layouts[layout])
    title_frame(slide)
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.text_frame.text = title
            for p in ph.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(20)
                    r.font.name = DISPLAY
                    r.font.color.rgb = WHITE
            break
    return slide


def text(slide, x, y, w, h, lines, *, size=12, colour=WHITE, bold=False,
         align=PP_ALIGN.LEFT, font=BODY, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(_in(x), _in(y), _in(w), _in(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = _in(0.04)
    tf.margin_top = tf.margin_bottom = _in(0.02)
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(3)
        parts = line if isinstance(line, list) else [(line, bold, colour, size)]
        for t, b, c, s in parts:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(s)
            r.font.bold = b
            r.font.name = font
            r.font.color.rgb = c
    return box


def card(slide, x, y, w, h, *, fill=None, line=ACCENT, width=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, _in(x), _in(y), _in(w), _in(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(width)
    s.shadow.inherit = False
    if s.has_text_frame:
        s.text_frame.text = ""
    return s


def bullets(slide, x, y, w, items, *, size=12.5, gap=0.42, marker=True):
    yy = y
    for item in items:
        if marker:
            d = card(slide, x, yy + 0.075, 0.105, 0.105, fill=ACCENT, line=None,
                     shape=MSO_SHAPE.OVAL)
        parts = item if isinstance(item, list) else [(item, False, WHITE, size)]
        text(slide, x + (0.24 if marker else 0), yy - 0.03, w - 0.24, gap + 0.2,
             [parts], size=size, spacing=0.95)
        yy += gap
    return yy


def picture(slide, name, x, y, w=None, h=None, *, frame=True):
    """Place a screenshot, preserving aspect ratio, inside an accent frame."""
    from PIL import Image
    path = SHOTS / f"{name}.png"
    iw, ih = Image.open(path).size
    if w and not h:
        h = w * ih / iw
    elif h and not w:
        w = h * iw / ih
    else:
        scale = min(w / iw, h / ih)
        w, h = iw * scale, ih * scale
    if frame:
        card(slide, x - 0.035, y - 0.035, w + 0.07, h + 0.07,
             fill=None, line=ACCENT, width=1.0, shape=MSO_SHAPE.RECTANGLE)
    slide.shapes.add_picture(str(path), _in(x), _in(y), _in(w), _in(h))
    return w, h


def chevron(slide, x, y, w, h, label, *, fill=None, line=ACCENT, size=10.5,
            colour=WHITE):
    s = card(slide, x, y, w, h, fill=fill, line=line)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = _in(0.03)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 0.92
    for i, part in enumerate(label.split("\n")):
        pp = p if i == 0 else tf.add_paragraph()
        pp.alignment = PP_ALIGN.CENTER
        pp.line_spacing = 0.92
        r = pp.add_run()
        r.text = part
        r.font.size = Pt(size)
        r.font.name = BODY
        r.font.color.rgb = colour
        r.font.bold = (i == 0)
    return s


def arrow(slide, x, y, w=0.26, h=0.14, right=True):
    s = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW if right else MSO_SHAPE.DOWN_ARROW,
        _in(x), _in(y), _in(w), _in(h))
    s.fill.solid()
    s.fill.fore_color.rgb = ACCENT
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def stat(slide, x, y, w, big, caption, *, colour=ACCENT, big_size=30):
    text(slide, x, y, w, 0.55, big, size=big_size, colour=colour, bold=True,
         align=PP_ALIGN.CENTER, font=DISPLAY)
    text(slide, x, y + 0.52, w, 0.6, caption.split("\n"), size=10,
         colour=MUTED, align=PP_ALIGN.CENTER, spacing=0.95)


def note(slide, y, message, *, colour=MUTED, size=9.5):
    text(slide, X0, y, W, 0.3, message, size=size, colour=colour,
         align=PP_ALIGN.LEFT)


# ----------------------------------------------------------------- the slides

def slide_01(prs):
    slide = prs.slides.add_slide(prs.slide_masters[0].slide_layouts[L_TITLE])
    for ph in slide.placeholders:
        idx = ph.placeholder_format.idx
        if idx == 0:
            ph.text_frame.text = "AI RESEARCH PAPER ASSISTANT"
            for p in ph.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(30)
                    r.font.name = DISPLAY
                    r.font.color.rgb = WHITE
        elif idx == 1:
            ph.text_frame.text = "ResearchForge"
            for p in ph.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(18)
                    r.font.name = DISPLAY
                    r.font.color.rgb = ACCENT
    text(slide, 1.78, 4.62, 6.44, 0.8, [
        [("BIT4543 Artificial Intelligence", False, WHITE, 12)],
        [("AI Topic: Document Intelligence", False, MUTED, 10.5)],
        [("researchforge.rukon.dev", False, ACCENT, 10.5)],
    ], align=PP_ALIGN.CENTER, spacing=1.05)


def slide_02(prs):
    s = add_slide(prs, "Background and motivation")
    bullets(s, X0, 1.62, 5.15, [
        "Researchers work with large volumes of academic literature",
        "Reading and understanding papers manually is slow",
        "Key findings are hard to extract quickly",
        "Research gaps are difficult to identify",
        "Literature reviews require comparison across sources",
        "AI can assist with document understanding",
    ], size=12.5, gap=0.46)
    card(s, 6.25, 1.62, 2.97, 2.26, fill=NAVY, line=ACCENT)
    text(s, 6.42, 1.80, 2.63, 0.5, "THE FIRST PASS", size=11,
         colour=ACCENT, bold=True, font=DISPLAY, align=PP_ALIGN.CENTER)
    text(s, 6.42, 2.26, 2.63, 1.9, [
        [("For every paper a researcher must decide:", False, WHITE, 11)],
        [("What does it claim?", True, WHITE, 11)],
        [("What does it leave open?", True, WHITE, 11)],
        [("How does it relate to the rest?", True, WHITE, 11)],
    ], spacing=1.15)
    note(s, 4.55, "ResearchForge assists this first pass, so effort is spent on the "
                  "papers that matter.")


def slide_03(prs):
    s = add_slide(prs, "Problem statement")
    items = [
        ("Slow analysis", "Reading and assessing each paper\ntakes significant time"),
        ("Hard to extract", "Key findings are buried in\nlong, dense documents"),
        ("Gaps are hidden", "Identifying what a paper leaves\nopen requires close reading"),
        ("Review effort", "Literature reviews need\ncomparison across sources"),
        ("Duplicate AI cost", "Re-analysing the same paper\nrepeats API calls"),
        ("No structure", "Researchers lack a structured\nassistance workflow"),
    ]
    cw, ch, gx, gy = 2.68, 1.26, 2.88, 1.44
    for i, (head, body) in enumerate(items):
        x = X0 + (i % 3) * gx
        y = 1.62 + (i // 3) * gy
        card(s, x, y, cw, ch, fill=NAVY, line=ACCENT)
        text(s, x + 0.14, y + 0.13, cw - 0.28, 0.3, head, size=12,
             colour=ACCENT, bold=True, font=DISPLAY)
        text(s, x + 0.14, y + 0.47, cw - 0.28, 0.72, body.split("\n"),
             size=10.5, colour=WHITE, spacing=1.0)
    note(s, 4.62, "A general chatbot will answer about a paper it has not read - an "
                  "unreliable summary is worse than none.")


def slide_04(prs):
    s = add_slide(prs, "Project objectives")
    left = ["Upload research papers", "Extract text from PDF papers",
            "Generate summaries and key findings", "Identify research gaps",
            "Generate literature review content"]
    right = ["Support cross-paper literature review",
             "Provide a usable assistant interface",
             "Secure accounts and private ownership",
             "Reduce duplicate AI processing (cache)",
             "Controlled AI routing and fallback"]
    for col, items, x in ((0, left, X0), (1, right, X0 + 4.32)):
        yy = 1.66
        for i, item in enumerate(items):
            n = i + 1 + col * 5
            card(s, x, yy, 0.42, 0.42, fill=None, line=ACCENT,
                 shape=MSO_SHAPE.SNIP_2_DIAG_RECTANGLE)
            text(s, x, yy + 0.06, 0.42, 0.3, f"{n:02d}", size=11.5,
                 colour=ACCENT, bold=True, font=DISPLAY, align=PP_ALIGN.CENTER)
            text(s, x + 0.56, yy + 0.06, 3.52, 0.42, item, size=12, spacing=0.95)
            yy += 0.62
    note(s, 4.92, "Objectives 8 to 10 are the engineering that makes the analysis "
                  "trustworthy, private and affordable.")


def slide_05(prs):
    s = add_slide(prs, "Proposed solution: ResearchForge")
    steps = ["Academic\npaper (PDF)", "Upload", "PDF text\nextraction",
             "AI analysis", "Structured\ninsights"]
    x = X0
    cw = 1.42
    for i, st in enumerate(steps):
        fill = TEAL if i == 3 else NAVY
        chevron(s, x, 1.54, cw, 0.74, st, fill=fill, size=10)
        if i < len(steps) - 1:
            arrow(s, x + cw + 0.05, 1.83, 0.22, 0.16)
        x += cw + 0.32
    outs = ["Summary and key findings", "Research gaps",
            "Literature review", "Research library"]
    x = X0
    for o in outs:
        chevron(s, x, 2.48, 1.96, 0.46, o, fill=None, line=ACCENT, size=10)
        x += 2.12
    text(s, X0, 3.14, 4.32, 1.14, [
        [("Cross-paper review", True, ACCENT, 12.5)],
        [("Several analysed papers from the user's own library can be reviewed "
          "together. The feature requires more than one paper.", False, WHITE, 11.5)],
    ], spacing=1.15)
    text(s, X0, 4.34, 4.32, 0.9, [
        [("Every generated section is constrained to the uploaded document. Where "
          "the paper does not support a section, the system reports insufficient "
          "evidence.", False, MUTED, 10)],
    ], spacing=1.05)
    picture(s, "03-dashboard-empty", 5.35, 3.06, w=3.87)

def slide_06(prs):
    s = add_slide(prs, "Methodology")
    steps = ["Requirements", "System\ndesign", "Implementation", "AI\nintegration",
             "Database and\nauthentication", "Testing", "Production\nverification"]
    # Seven boxes must span exactly the 8.43 in content band, or the row runs
    # past the right edge of the title frame above it.
    x = X0
    cw, gap = 1.067, 0.16
    for i, st in enumerate(steps):
        chevron(s, x, 1.66, cw, 0.82, st, fill=NAVY, size=8.5)
        if i < len(steps) - 1:
            arrow(s, x + cw + 0.01, 1.99, 0.14, 0.14)
        x += cw + gap
    text(s, X0, 2.78, W, 0.4, "Incremental, milestone-based development",
         size=13, colour=ACCENT, bold=True, font=DISPLAY)
    bullets(s, X0, 3.20, W, [
        "Thirteen milestones, each ending with a test and a review before the next began",
        "Design decisions driven by measurement, not assumption",
        "Retrieval was not built after measuring that whole papers fit the model context",
        "Isolation moved into the database after a defect showed application filtering had failed silently",
        "Six additive-only migrations: no DROP, TRUNCATE or DELETE",
    ], size=11.5, gap=0.38)


def slide_07(prs):
    s = add_slide(prs, "System architecture")
    row1 = ["User", "Web interface\n(Next.js / React)",
            "FastAPI backend\nauthentication check",
            "PDF extraction,\nnormalisation, SHA-256 hash"]
    x = X0
    for i, label in enumerate(row1):
        chevron(s, x, 1.50, 1.95, 0.58, label, fill=NAVY, size=9.5)
        if i < len(row1) - 1:
            arrow(s, x + 1.98, 1.72, 0.18, 0.14)
        x += 2.16
    # Elbow connector: the pipeline ends at the right of row 1 and continues
    # at the left of row 2, so a single centred arrow would imply the wrong
    # source. Down from PDF extraction, left along the gap, into the cache.
    arrow(s, 8.16, 2.04, 0.14, 0.12, right=False)
    card(s, 1.52, 2.14, 6.78, 0.02, fill=ACCENT, line=None,
         shape=MSO_SHAPE.RECTANGLE)
    arrow(s, 1.52, 2.16, 0.14, 0.26, right=False)

    chevron(s, X0, 2.42, 1.62, 0.56, "Analysis cache\nlookup", fill=TEAL, size=9.5)
    arrow(s, 2.46, 2.63, 0.18, 0.14)
    chevron(s, 2.70, 2.42, 1.52, 0.56, "AI provider\nrouter", fill=NAVY, size=9.5)
    arrow(s, 4.26, 2.63, 0.18, 0.14)
    chevron(s, 4.50, 2.20, 2.00, 0.48, "Primary: Groq\nqwen/qwen3.6-27b",
            fill=None, line=ACCENT, size=9)
    chevron(s, 4.50, 2.76, 2.00, 0.48, "Fallback: Anthropic\nclaude-opus-5",
            fill=None, line=TEAL, size=9)
    arrow(s, 6.54, 2.63, 0.18, 0.14)
    chevron(s, 6.78, 2.42, 1.44, 0.56, "Schema\nvalidation", fill=NAVY, size=9.5)

    arrow(s, 7.43, 3.02, 0.14, 0.18, right=False)
    chevron(s, 6.30, 3.28, 2.92, 0.62,
            "Supabase PostgreSQL - Row Level Security", fill=TEAL, size=9.5)
    text(s, 6.30, 3.94, 2.92, 0.3, "Results filtered by the signed-in user",
         size=9, colour=MUTED, align=PP_ALIGN.CENTER)

    text(s, X0, 3.30, 5.30, 1.0, [
        [("A cache hit returns the stored analysis with no model call. A miss "
          "goes to the primary provider, with one fallback attempt permitted "
          "for an allowed retryable failure.", False, WHITE, 10.5)],
    ], spacing=1.1)

    card(s, X0, 4.48, W, 0.66, fill=None, line=MUTED, width=0.75)
    text(s, X0 + 0.14, 4.56, W - 0.28, 0.52, [
        [("Prepared scaffolding, not production: ", True, MUTED, 9.5),
         ("Jina embeddings, pgvector and the chunks table exist in the "
          "repository but are not called during analysis.", False, MUTED, 9.5)],
        [("The pipeline is full-document grounded generation, not RAG.",
          False, MUTED, 9.5)],
    ], spacing=1.05)

def slide_08(prs):
    s = add_slide(prs, "AI model implementation")
    text(s, X0, 1.56, 4.9, 0.3, "INFERENCE, NOT TRAINING", size=11,
         colour=ACCENT, bold=True, font=DISPLAY)
    bullets(s, X0, 1.94, 4.9, [
        [("Primary: ", True, WHITE, 11.5), ("Groq ", False, WHITE, 11.5),
         ("qwen/qwen3.6-27b", False, ACCENT, 11.5)],
        [("Fallback: ", True, WHITE, 11.5), ("Anthropic ", False, WHITE, 11.5),
         ("claude-opus-5", False, ACCENT, 11.5)],
        "Three structured passes: summary, gaps, review",
        "Whole uploaded document supplied as grounding context",
        "Output validated against a declared schema",
        "Invalid output is discarded, never repaired",
        "Provenance recorded: provider, model, fallback, cache",
    ], size=11.5, gap=0.4)
    card(s, 5.95, 1.52, 3.27, 2.02, fill=NAVY, line=TEAL)
    text(s, 6.10, 1.60, 2.97, 0.3, "CONTROLLED ROUTING", size=10.5,
         colour=TEAL, bold=True, font=DISPLAY)
    text(s, 6.10, 1.92, 2.97, 1.56, [
        [("The owner selects one global primary provider; the other enabled "
          "provider becomes the fallback.", False, WHITE, 10)],
        [("Normal users cannot switch providers. There is no auto option.",
          False, WHITE, 10)],
        [("One fallback attempt only, for allowed retryable failures.",
          True, WHITE, 10)],
    ], spacing=1.05)
    picture(s, "10-settings-ai", 5.95, 3.66, w=3.27)


def slide_09(prs):
    s = add_slide(prs, "Key system features")
    w, h = picture(s, "05-summary", X0, 1.56, w=4.36)
    text(s, X0, 1.56 + h + 0.12, 4.36, 0.3,
         "One upload produces summary, key findings, gaps and review",
         size=10, colour=MUTED, align=PP_ALIGN.CENTER)
    bullets(s, 5.42, 1.60, 3.80, [
        "PDF upload and text extraction",
        "Scanned documents rejected, not analysed",
        "Structured summary and key findings",
        "Research gaps and stated limitations",
        "Literature review content",
        "Cross-paper review across saved papers",
        "Private per-user research library",
    ], size=11.5, gap=0.42)
    text(s, X0, 4.46, 4.36, 0.7,
         "Each section is generated from the uploaded document only, and validated "
         "against a declared schema before it is stored.",
         size=10, colour=MUTED, spacing=1.05)

def slide_10(prs):
    s = add_slide(prs, "Research gaps and literature review")
    w, h = picture(s, "06-gaps", X0, 1.60, w=3.86)
    text(s, X0, 1.60 + h + 0.10, 3.86, 0.3, "Research gaps and limitations",
         size=11, colour=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    w2, h2 = picture(s, "07-review", 5.36, 1.60, w=3.86)
    text(s, 5.36, 1.60 + h2 + 0.10, 3.86, 0.3, "Cross-paper literature review",
         size=11, colour=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    bullets(s, X0, 4.18, W, [
        "Gaps are drawn from the limitations the authors state and the future work they propose",
        "Literature review content covers themes, comparisons with prior work and future directions",
        "Cross-paper review compares several analysed papers; it requires more than one paper",
        "Where the paper does not support a section, the system reports insufficient evidence",
    ], size=11, gap=0.30)


def slide_11(prs):
    s = add_slide(prs, "Authentication and data ownership")
    bullets(s, X0, 1.60, 4.55, [
        "Email and password sign-in, plus Google sign-in",
        "Password reset and account settings",
        "Private per-user research library",
        "Ownership enforced by PostgreSQL Row Level Security",
        "Owner rights resolved server-side, not in the browser",
        "No password is stored by this project",
    ], size=11.5, gap=0.4)
    card(s, X0, 4.12, 4.55, 1.02, fill=NAVY, line=ACCENT)
    text(s, X0 + 0.14, 4.22, 4.27, 0.84, [
        [("User A cannot access User B's private records.", True, ACCENT, 11)],
        [("Verified live: 26 isolation checks, 0 failures, including a direct "
          "database query that bypassed the application and returned zero rows.",
          False, WHITE, 10)],
    ], spacing=1.05)
    picture(s, "08-library", 5.62, 1.60, w=3.60)
    picture(s, "09-settings-account", 5.62, 3.62, w=3.60)


def slide_12(prs):
    s = add_slide(prs, "Same-paper analysis cache")
    chain = ["PDF upload", "Text extraction", "Normalisation", "SHA-256 hash"]
    x = X0
    for i, c in enumerate(chain):
        chevron(s, x, 1.52, 1.72, 0.46, c, fill=NAVY, size=10)
        if i < len(chain) - 1:
            arrow(s, x + 1.75, 1.68, 0.2, 0.14)
        x += 1.95
    arrow(s, 4.93, 2.04, 0.14, 0.16, right=False)
    chevron(s, 3.72, 2.26, 2.55, 0.46, "Cache lookup", fill=TEAL, size=10.5)

    card(s, X0, 2.92, 4.16, 1.46, fill=None, line=ACCENT)
    text(s, X0 + 0.12, 2.99, 3.9, 0.28, "HIT - existing valid analysis", size=10,
         colour=ACCENT, bold=True, font=DISPLAY)
    text(s, X0 + 0.12, 3.27, 3.9, 1.06, [
        [("Reuse the stored analysis", False, WHITE, 10)],
        [("Create this user's own private record", False, WHITE, 10)],
        [("Display result - 0 AI calls", True, WHITE, 10)],
        [("Observed: 3.2 s same user, 2.7 s another user", False, ACCENT, 10)],
    ], spacing=1.05)

    card(s, 5.06, 2.92, 4.16, 1.46, fill=None, line=TEAL)
    text(s, 5.18, 2.99, 3.9, 0.28, "MISS - no stored analysis", size=10,
         colour=TEAL, bold=True, font=DISPLAY)
    text(s, 5.18, 3.27, 3.9, 1.06, [
        [("Run AI analysis and validate the response", False, WHITE, 10)],
        [("Store the analysis and the cache entry", False, WHITE, 10)],
        [("Only successful analyses are cached", True, WHITE, 10)],
        [("Observed: 32.4 s with three AI calls", False, TEAL, 10)],
    ], spacing=1.05)

    card(s, X0, 4.52, W, 0.62, fill=NAVY, line=None)
    text(s, X0 + 0.14, 4.60, W - 0.28, 0.48, [
        [("The cache is global internal infrastructure, not a shared library.",
          True, WHITE, 10)],
        [("Each user keeps their own private record, the original uploader is "
          "never revealed, and Row Level Security still applies.", False, MUTED, 10)],
    ], spacing=1.05)

def slide_13(prs):
    s = add_slide(prs, "Testing and evaluation")
    stats = [("522", "backend tests\npassing"), ("16", "frontend tests\npassing"),
             ("26", "live isolation\nchecks passed"), ("0", "failing tests")]
    x = X0
    for big, cap in stats:
        stat(s, x, 1.54, 1.98, big, cap)
        x += 2.12
    text(s, X0, 2.76, W, 0.3, "VERIFIED ON THE CURRENT REPOSITORY", size=10.5,
         colour=ACCENT, bold=True, font=DISPLAY)
    left = ["Unit testing against fakes, so the suite runs offline",
            "Integration and API testing on the deployed system",
            "Authentication and session testing",
            "User ownership and Row Level Security testing"]
    right = ["AI provider and fallback testing",
             "Cache testing with four independent signals",
             "Responsive testing on a real device",
             "Ruff, TypeScript and production build all clean"]
    bullets(s, X0, 3.16, 4.20, left, size=11, gap=0.40)
    bullets(s, X0 + 4.32, 3.16, 4.10, right, size=11, gap=0.40)
    note(s, 4.86, "5 further tests call live providers and are skipped by design. "
                  "No model accuracy value is claimed: there is no labelled benchmark.")

def slide_14(prs):
    s = add_slide(prs, "Results and achievements")
    text(s, X0, 1.54, W, 0.3, "OBSERVED IN THE TESTED SCENARIOS", size=10.5,
         colour=ACCENT, bold=True, font=DISPLAY)
    rows = [("Cache miss", "32.4 s", "3 AI calls"),
            ("Cache hit, same user", "3.2 s", "0 AI calls"),
            ("Cache hit, different user", "2.7 s", "0 AI calls")]
    y = 1.92
    for label, t, calls in rows:
        card(s, X0, y, 4.3, 0.46, fill=NAVY, line=None)
        text(s, X0 + 0.14, y + 0.09, 2.2, 0.3, label, size=10.5)
        text(s, X0 + 2.34, y + 0.07, 0.9, 0.3, t, size=12, colour=ACCENT,
             bold=True, align=PP_ALIGN.RIGHT)
        text(s, X0 + 3.32, y + 0.09, 0.86, 0.3, calls, size=10,
             colour=MUTED, align=PP_ALIGN.RIGHT)
        y += 0.54
    text(s, X0, 3.60, 4.3, 0.6, [
        [("A repeated analysis costs about a tenth of the time and no tokens, "
          "while each user keeps a separate private record.", False, WHITE, 10.5)],
    ], spacing=1.05)

    card(s, 5.36, 1.88, 3.86, 2.34, fill=None, line=TEAL)
    text(s, 5.50, 1.98, 3.58, 0.3, "PROVIDER FINDING", size=10.5,
         colour=TEAL, bold=True, font=DISPLAY)
    text(s, 5.50, 2.30, 3.58, 1.84, [
        [("Under the owner's configuration all five evaluation papers completed, "
          "but every one was produced by the fallback provider.", False, WHITE, 10)],
        [("At its free service tier the primary provider allows 7,000 input "
          "tokens per minute, while the papers required 7,920 to 20,362 in a "
          "single request.", False, WHITE, 10)],
        [("The fallback worked correctly - and in doing so concealed that the "
          "primary was unusable.", True, ACCENT, 10)],
    ], spacing=1.05)
    note(s, 4.42, "Isolation verified live: 26 checks, 0 failures. Deployed and "
                  "publicly reachable. No universal performance claim is made.")


def slide_15(prs):
    s = add_slide(prs, "Conclusion and future work")
    text(s, X0, 1.54, W, 0.66, [
        [("ResearchForge provides an AI-assisted workflow for analysing academic "
          "research papers and extracting structured research insights.",
          False, WHITE, 12.5)],
    ], spacing=1.1)
    card(s, X0, 2.24, 4.16, 2.44, fill=NAVY, line=ACCENT)
    text(s, X0 + 0.16, 2.34, 3.84, 0.3, "DELIVERED", size=11,
         colour=ACCENT, bold=True, font=DISPLAY)
    bullets(s, X0 + 0.16, 2.70, 3.84, [
        "Upload, extraction and AI analysis",
        "Research gaps and literature review",
        "Cross-paper review",
        "Authentication and per-user ownership",
        "Provider routing with one fallback",
        "Same-paper analysis cache",
    ], size=10.5, gap=0.31)

    card(s, 5.06, 2.24, 4.16, 2.44, fill=None, line=TEAL)
    text(s, 5.22, 2.34, 3.84, 0.3, "FUTURE WORK", size=11,
         colour=TEAL, bold=True, font=DISPLAY)
    bullets(s, 5.22, 2.70, 3.84, [
        "Production vector retrieval (RAG)",
        "Stronger retrieval and ranking",
        "Citation-aware evidence linking",
        "Larger benchmark and human evaluation",
        "Document versioning and multilingual support",
        "Richer research comparison",
    ], size=10.5, gap=0.31)
    text(s, X0, 4.82, W, 0.4, "researchforge.rukon.dev", size=12,
         colour=ACCENT, bold=True, font=DISPLAY, align=PP_ALIGN.CENTER)


def build():
    prs = Presentation(str(TEMPLATE))
    clear_slides(prs)
    for fn in (slide_01, slide_02, slide_03, slide_04, slide_05, slide_06,
               slide_07, slide_08, slide_09, slide_10, slide_11, slide_12,
               slide_13, slide_14, slide_15):
        fn(prs)
    OUT.mkdir(exist_ok=True)
    path = OUT / "ResearchForge_Final_Presentation.pptx"
    prs.save(path)
    return path, len(prs.slides.__iter__.__self__._sldIdLst)


if __name__ == "__main__":
    p, n = build()
    print(f"  {p.name}")
    print(f"  {p.stat().st_size:,} bytes")
    print(f"  slides: {n}")
